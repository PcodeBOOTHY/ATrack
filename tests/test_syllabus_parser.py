"""Syllabus parser tests with a fake Anthropic client (no network, no API key)."""

import base64
import io
import json
from datetime import date
from types import SimpleNamespace

import pytest
from docx import Document

from core.extraction import ExtractionError
from services import syllabus_parser as sp

VALID = json.dumps({"courses": [{
    "course": {"code": "ME 214", "name": None, "section": None, "instructor": None, "term": "Winter 2026"},
    "assessments": [{
        "title": "Midterm", "type": "midterm", "weight_percent": 25, "due_date": "2026-02-26",
        "due_time": "19:00", "location": "ENG 1B71", "notes": None, "confidence": 0.9,
        "source_quote": "Midterm (25%) Feb 26, 7 pm",
    }],
}]})


class FakeStream:
    def __init__(self, message):
        self.message = message

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def get_final_message(self):
        return self.message


class FakeClient:
    """Mimics client.beta.messages.stream(...) and records every request."""

    def __init__(self, *replies):
        self.replies = list(replies)
        self.calls = []
        self.beta = SimpleNamespace(messages=SimpleNamespace(stream=self._stream))

    def _stream(self, **kwargs):
        self.calls.append(kwargs)
        reply = self.replies.pop(0)
        if isinstance(reply, Exception):
            raise reply
        text, stop = reply if isinstance(reply, tuple) else (reply, "end_turn")
        content = [SimpleNamespace(type="thinking", thinking=""), SimpleNamespace(type="text", text=text)]
        return FakeStream(SimpleNamespace(stop_reason=stop, content=content))


def _docx_bytes():
    doc = Document()
    doc.add_paragraph("ME 214 Dynamics - Winter 2026")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Assessment", "Weight"
    table.cell(1, 0).text, table.cell(1, 1).text = "Midterm", "25%"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


TODAY = date(2026, 1, 5)


# --- content building -------------------------------------------------------

def test_pdf_and_image_sent_as_blocks():
    files = [sp.UploadedDoc("outline.pdf", b"%PDF-1.4"), sp.UploadedDoc("shot.JPG", b"\xff\xd8")]
    blocks = sp.build_content(files, None, TODAY)
    assert [b["type"] for b in blocks] == ["document", "image", "text"]
    assert blocks[0]["source"]["media_type"] == "application/pdf"
    assert base64.b64decode(blocks[0]["source"]["data"]) == b"%PDF-1.4"
    assert blocks[1]["source"]["media_type"] == "image/jpeg"
    assert "2026-01-05" in blocks[-1]["text"]


def test_docx_converted_to_text_including_tables():
    text = sp.docx_to_text(_docx_bytes())
    assert "ME 214 Dynamics" in text
    assert "Midterm | 25%" in text
    blocks = sp.build_content([sp.UploadedDoc("syllabus.docx", _docx_bytes())], None, TODAY)
    assert blocks[0]["type"] == "text" and "Midterm | 25%" in blocks[0]["text"]


def test_pasted_text_included():
    blocks = sp.build_content([], "Quiz 1 worth 5%", TODAY)
    assert "Quiz 1 worth 5%" in blocks[0]["text"]


@pytest.mark.parametrize("files,text,message", [
    ([], "   ", "Upload at least one file"),
    ([sp.UploadedDoc("old.doc", b"x")], None, "unsupported file type"),
    ([sp.UploadedDoc("big.png", b"0" * (sp.MAX_IMAGE_BYTES + 1))], None, "larger than 5 MB"),
    ([sp.UploadedDoc("bad.docx", b"not a zip")], None, "Could not read"),
])
def test_bad_input_rejected(files, text, message):
    with pytest.raises(sp.UnsupportedFileError, match=message):
        sp.build_content(files, text, TODAY)


# --- requests ---------------------------------------------------------------

def test_successful_extraction_and_request_shape():
    client = FakeClient(VALID)
    result = sp.SyllabusParser(client, model="claude-sonnet-5-5").extract([], "Midterm 25%", TODAY)
    assert result.courses[0].assessments[0].title == "Midterm"
    call = client.calls[0]
    assert call["model"] == "claude-sonnet-5-5"
    assert call["output_config"]["format"]["type"] == "json_schema"
    assert call["fallbacks"] == "default" and call["betas"] == [sp.FALLBACK_BETA]


def test_no_fallbacks_for_other_models():
    client = FakeClient(VALID)
    sp.SyllabusParser(client, model="claude-haiku-5-5").extract([], "x", TODAY)
    assert "fallbacks" not in client.calls[0] and "betas" not in client.calls[0]


def test_retries_once_on_invalid_output():
    client = FakeClient('{"courses": "nope"}', VALID)
    result = sp.SyllabusParser(client, model="m").extract([], "x", TODAY)
    assert len(result.courses) == 1
    assert len(client.calls) == 2
    retry_text = client.calls[1]["messages"][0]["content"][-1]["text"]
    assert "invalid output" in retry_text


def test_gives_up_after_second_invalid_output():
    client = FakeClient("not json", "still not json")
    with pytest.raises(ExtractionError):
        sp.SyllabusParser(client, model="m").extract([], "x", TODAY)
    assert len(client.calls) == 2


@pytest.mark.parametrize("stop,message", [("refusal", "declined"), ("max_tokens", "too much output")])
def test_bad_stop_reasons(stop, message):
    client = FakeClient(("", stop))
    with pytest.raises(ExtractionError, match=message):
        sp.SyllabusParser(client, model="m").extract([], "x", TODAY)


def test_connection_error_is_friendly():
    import anthropic
    import httpx2

    err = anthropic.APIConnectionError(request=httpx2.Request("POST", "https://api.anthropic.com"))
    with pytest.raises(ExtractionError, match="internet"):
        sp.SyllabusParser(FakeClient(err), model="m").extract([], "x", TODAY)
